"""Carga masiva segura de sucursales; el módulo no accede por su cuenta a SQL Server.

El backend autentica, valida la empresa, revalida los registros contra SQL y ejecuta
SPs existentes dentro de una única transacción.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

FIELDS = ('codigo', 'nombre', 'telefono', 'direccion', 'municipio', 'departamento', 'referencia_llegada', 'observaciones')
REQUIRED = ('nombre', 'direccion', 'municipio')
LIMITS = {'codigo': 30, 'nombre': 180, 'telefono': 20, 'direccion': 450, 'municipio': 100,
          'departamento': 100, 'referencia_llegada': 400, 'observaciones': 500}
MAX_ROWS = 300
MAX_BYTES = 4 * 1024 * 1024
ALIASES = {
    'codigo': 'codigo', 'codigo_sucursal': 'codigo', 'codigo_de_sucursal': 'codigo',
    'nombre': 'nombre', 'sucursal': 'nombre', 'nombre_sucursal': 'nombre', 'nombre_de_sucursal': 'nombre',
    'telefono': 'telefono', 'telefono_sucursal': 'telefono',
    'direccion': 'direccion', 'direccion_completa': 'direccion',
    'municipio': 'municipio', 'departamento': 'departamento',
    'referencia': 'referencia_llegada', 'referencia_llegada': 'referencia_llegada',
    'referencia_para_llegar': 'referencia_llegada', 'observaciones': 'observaciones',
}


def _key(value):
    s = unicodedata.normalize('NFKD', str(value or ''))
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', '_', s.lower()).strip('_')


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _csv_rows(data):
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError('No se pudo interpretar el CSV. Exportalo como UTF-8.')
    try:
        dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(decoded, newline=''), dialect=dialect))


def _xml_rows(data):
    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
        raise ValueError('XML con entidades o DTD no está permitido.')
    try:
        root = ET.fromstring(data)
    except ET.ParseError as exc:
        raise ValueError('El XML no tiene una estructura válida.') from exc
    if _key(root.tag.split('}')[-1]) != 'sucursales':
        raise ValueError('La raíz del XML debe ser <sucursales>.')
    records = []
    for node in list(root):
        if _key(node.tag.split('}')[-1]) != 'sucursal':
            raise ValueError('Dentro de <sucursales> solamente se permiten <sucursal>.')
        records.append({child.tag.split('}')[-1]: (child.text or '') for child in node})
    return records


def _xlsx_rows(data):
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError('El archivo .xlsx no es un libro de Excel válido.')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 12 * 1024 * 1024:
            raise ValueError('El Excel descomprimido supera el tamaño permitido.')
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError('Falta instalar openpyxl: instalá las dependencias de requirements.txt.') from exc
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book.active
            rows = []
            for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if idx > MAX_ROWS + 12:
                    raise ValueError('El Excel tiene demasiadas filas.')
                rows.append([_text(cell) for cell in row])
            return rows
        finally:
            book.close()
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        raise ValueError(f'No se pudo leer el Excel: {str(exc)[:120]}') from exc


def parse_upload(upload):
    if not upload or not upload.filename:
        raise ValueError('Seleccioná un archivo .xlsx, .csv o .xml.')
    ext = Path(upload.filename).suffix.lower()
    if ext not in {'.xlsx', '.csv', '.xml'}:
        raise ValueError('Formato no admitido. Utilizá .xlsx, .csv o .xml.')
    data = upload.stream.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ValueError('El archivo supera 4 MB.')
    if not data:
        raise ValueError('El archivo está vacío.')
    raw = _csv_rows(data) if ext == '.csv' else _xml_rows(data) if ext == '.xml' else _xlsx_rows(data)
    if not raw:
        raise ValueError('No se encontraron sucursales.')
    if isinstance(raw[0], dict):
        pairs = [(i + 1, row) for i, row in enumerate(raw)]
    else:
        header_at = next((i for i, r in enumerate(raw[:10]) if 'nombre' in {_key(v).replace('nombre_de_sucursal', 'nombre') for v in r}
                          or any(_key(v) in ('nombre_sucursal', 'sucursal') for v in r)), None)
        if header_at is None:
            raise ValueError('No se encontró el encabezado Nombre. Usá la plantilla de SEPRIGUA.')
        headers = [ALIASES.get(_key(x), _key(x)) for x in raw[header_at]]
        if len(set(headers)) != len(headers):
            raise ValueError('El archivo tiene columnas repetidas.')
        pairs = [(idx + 1, dict(zip(headers, values))) for idx, values in enumerate(raw[header_at + 1:], header_at + 1)]
    if len(pairs) > MAX_ROWS:
        raise ValueError(f'Por seguridad se permiten hasta {MAX_ROWS} sucursales por archivo.')
    result = []
    for number, raw_row in pairs:
        clean = {ALIASES.get(_key(k), _key(k)): _text(v) for k, v in raw_row.items()}
        if not any(clean.values()):
            continue
        item = {field: clean.get(field, '') for field in FIELDS}
        item['departamento'] = item['departamento'] or 'Guatemala'
        mistakes = [f'{field} es obligatorio' for field in REQUIRED if not item[field]]
        mistakes += [f'{field} supera {LIMITS[field]} caracteres' for field in FIELDS if len(item[field]) > LIMITS[field]]
        item['fila'] = number
        item['errores'] = mistakes
        result.append(item)
    if not result:
        raise ValueError('El archivo no contiene sucursales con datos.')
    return result


def validate_rows(items, existing):
    """Devuelve filas con estados NUEVA, DUPLICADA o ERROR sin mutar el origen."""
    taken_codes = {_key(x.get('codigo')) for x in existing if x.get('codigo')}
    taken_names = {_key(x.get('nombre')) for x in existing if x.get('nombre')}
    seen_codes = set()
    seen_names = set()
    result = []
    for row in items:
        record = dict(row)
        code = _key(record.get('codigo'))
        name = _key(record.get('nombre'))
        if record['errores']:
            record['estado'] = 'ERROR'
        elif (code and (code in seen_codes or code in taken_codes)) or name in seen_names or name in taken_names:
            record['estado'] = 'DUPLICADA'
        else:
            record['estado'] = 'NUEVA'
        if code:
            seen_codes.add(code)
        seen_names.add(name)
        result.append(record)
    return result
